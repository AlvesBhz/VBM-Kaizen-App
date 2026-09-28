#!/usr/bin/env python3
# =====================================================================
#  Gerador dos scripts carga_hist_importe_<ini>_<fim>.sql
#  ---------------------------------------------------------------------
#  Le APENAS as guias que comecam com KZN_HIST da planilha de extracao e
#  emite um script T-SQL de carga com valores literais (sem dependencia
#  do arquivo no momento da execucao).
#
#      python3 gerar_carga_hist.py <planilha.xlsx> <id_ini> <id_fim> <saida.sql>
#
#  A lista de Kaizens vem SEMPRE da guia KZN_HIST_PEDRAVISAOCONSOLIDADA;
#  as auxiliares sao filtradas por essa lista, nao pela faixa, porque os
#  IDs nao sao contiguos.
#
#  Versionado de proposito: este gerador ja foi perdido junto com o
#  container uma vez, e reconstruir os literais na mao e caro.
# =====================================================================
import sys, datetime, decimal
import openpyxl

GUIA_PVC = 'KZN_HIST_PEDRAVISAOCONSOLIDADA'
AUX = [
    # (guia, tabela destino, colunas, lote)
    ('KZN_HIST_MEMBROS_EQUIPE',     'KZN_HIST_MEMBROS_EQUIPE',
     ['ID_KAIZEN', 'ID_USUARIO', 'DT_ATUALIZACAO'],
     'ID_KAIZEN INT, ID_USUARIO INT, DT_ATUALIZACAO DATETIME2(3)', 100),
    ('KZN_HIST_RESULTADO_KAIZEN',   'KZN_HIST_RESULTADO_KAIZEN',
     ['ID_KAIZEN', 'ID_RESULTADO', 'DT_ATUALIZACAO'],
     'ID_KAIZEN INT, ID_RESULTADO INT, DT_ATUALIZACAO DATETIME2(3)', 100),
    ('KZN_HIST_KAIZEN_HIERARQUIA',  'KZN_HIST_KAIZEN_HIERARQUIA',
     ['ID_KAIZEN', 'ID_USUARIO_LIDER'] + ['NM_HIERARQUIA_N%d' % i for i in range(1, 9)] + ['DT_ATUALIZACAO'],
     'ID_KAIZEN INT, ID_USUARIO_LIDER INT, ' +
     ', '.join('NM_HIERARQUIA_N%d NVARCHAR(MAX)' % i for i in range(1, 9)) +
     ', DT_ATUALIZACAO DATETIME2(3)', 50),
    ('KZN_HIST_KAIZEN_DESPERDICIO', 'KZN_HIST_KAIZEN_DESPERDICIO',
     ['ID_KAIZEN', 'ID_DESPERDICIO', 'DT_ATUALIZACAO'],
     'ID_KAIZEN INT, ID_DESPERDICIO INT, DT_ATUALIZACAO DATETIME2(3)', 100),
]

# Colunas da guia principal, na ordem, e como cada uma e emitida.
# 'txt' = literal string sempre; 'int'/'num' = numerico; 'dt' = datetime.
# Os quatro IDs que trazem '#N/A' sao 'txt' de proposito: entram no staging
# como texto e sao convertidos por TRY_CONVERT na E2.
PVC_COLS = [
    ('ID_KAIZEN',               'ID_KAIZEN',               'int'),
    ('ID_USUARIO_CADASTRO',     'TX_CADASTRO',             'txt'),
    ('ID_USUARIO_LIDER',        'TX_LIDER',                'txt'),
    ('NM_KAIZEN',               'NM_KAIZEN',               'txt'),
    ('ID_CATEGORIA',            'ID_CATEGORIA',            'int'),
    ('ID_REPLICACAO',           'TX_REPLICACAO',           'txt'),
    ('DS_PROBLEMA',             'DS_PROBLEMA',             'txt'),
    ('DS_OBJETIVO',             'DS_OBJETIVO',             'txt'),
    ('ID_APROVADOR',            'TX_APROVADOR',            'txt'),
    ('URL_IMG_ANTES',           'URL_IMG_ANTES',           'txt'),
    ('DS_ESTADO_ANTES',         'DS_ESTADO_ANTES',         'txt'),
    ('URL_IMG_DEPOIS',          'URL_IMG_DEPOIS',          'txt'),
    ('DS_ESTADO_DEPOIS',        'DS_ESTADO_DEPOIS',        'txt'),
    ('URL_REFERENCIA',          'URL_REFERENCIA',          'txt'),
    ('DS_COMPARA_META',         'DS_COMPARA_META',         'txt'),
    ('DS_LICOES_APRENDIDAS',    'DS_LICOES_APRENDIDAS',    'txt'),
    ('VL_RESULTADO_FINANCEIRO', 'VL_RESULTADO_FINANCEIRO', 'num'),
    ('ID_MOEDA',                'ID_MOEDA',                'int'),
    ('DS_RESULTADO_ALCANCADO',  'DS_RESULTADO_ALCANCADO',  'txt'),
    ('DT_CONCLUSAO',            'DT_CONCLUSAO',            'dt'),
    ('DT_ATUALIZACAO',          'DT_REG',                  'dt'),
    ('ID_USUARIO_ATUALIZACAO',  'TX_ATUALIZACAO',          'txt'),
    ('DS_MOTIVO',               'DS_MOTIVO',               'txt'),
    ('ID_STATUS',               'ID_STATUS',               'int'),
]

STAGING_PVC = """    ID_KAIZEN INT, TX_CADASTRO NVARCHAR(50), TX_LIDER NVARCHAR(50), NM_KAIZEN NVARCHAR(MAX),
    ID_CATEGORIA INT, TX_REPLICACAO NVARCHAR(200), DS_PROBLEMA NVARCHAR(MAX), DS_OBJETIVO NVARCHAR(MAX),
    TX_APROVADOR NVARCHAR(50), URL_IMG_ANTES NVARCHAR(MAX), DS_ESTADO_ANTES NVARCHAR(MAX), URL_IMG_DEPOIS NVARCHAR(MAX),
    DS_ESTADO_DEPOIS NVARCHAR(MAX), URL_REFERENCIA NVARCHAR(MAX), DS_COMPARA_META NVARCHAR(MAX),
    DS_LICOES_APRENDIDAS NVARCHAR(MAX), VL_RESULTADO_FINANCEIRO DECIMAL(18,2), ID_MOEDA INT,
    DS_RESULTADO_ALCANCADO NVARCHAR(MAX), DT_CONCLUSAO DATE, DT_REG DATETIME2(3),
    TX_ATUALIZACAO NVARCHAR(50), DS_MOTIVO NVARCHAR(MAX), ID_STATUS INT"""

# Colunas da HIST principal que esta carga preenche (sem a de data, que e
# resolvida em tempo de execucao).
PVC_DESTINO = [c[0] for c in PVC_COLS if c[0] != 'DT_ATUALIZACAO']


def lit_txt(v):
    if v is None:
        return 'NULL'
    s = str(v)
    if s == '':
        return 'NULL'
    return "N'" + s.replace("'", "''") + "'"


def lit_num(v):
    if v is None:
        return 'NULL'
    if isinstance(v, decimal.Decimal):
        return str(v)
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def lit_dt(v):
    if v is None:
        return 'NULL'
    if isinstance(v, datetime.date) and not isinstance(v, datetime.datetime):
        v = datetime.datetime(v.year, v.month, v.day)
    if not isinstance(v, datetime.datetime):
        return lit_txt(v)
    # o fechamento da quote DEPOIS dos milissegundos: '%s' % x + '%03d' % ms
    # colocaria a quote no meio e gerava "Msg 102 near '000'"
    return "'%s%03d'" % (v.strftime('%Y-%m-%dT%H:%M:%S.'), v.microsecond // 1000)


assert lit_dt(datetime.datetime(2025, 12, 24)) == "'2025-12-24T00:00:00.000'", 'auto-teste do literal de data falhou'


def lit(v, kind):
    return {'txt': lit_txt, 'int': lit_num, 'num': lit_num, 'dt': lit_dt}[kind](v)


def ler(wb, guia, ncols, skip):
    ws = wb[guia]
    it = ws.iter_rows(values_only=True)
    for _ in range(skip):
        next(it)
    return [r[:ncols] for r in it if r[0] is not None]


def maxlen(rows, idx):
    return max((len(str(r[idx])) for r in rows if isinstance(r[idx], str)), default=0)


def main():
    src, ini, fim, out = sys.argv[1], int(sys.argv[2]), int(sys.argv[3]), sys.argv[4]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    faltando = [g for g, *_ in [(GUIA_PVC,)] + [(a[0],) for a in AUX] if g not in wb.sheetnames]
    if faltando:
        sys.exit('guia(s) ausente(s) na planilha: %s' % ', '.join(faltando))

    # A guia principal tem 1 linha de anotacao acima do cabecalho; as
    # auxiliares nao. Confirmado na planilha de origem.
    pvc = [r for r in ler(wb, GUIA_PVC, 24, 2) if ini <= int(r[0]) <= fim]
    pvc.sort(key=lambda r: int(r[0]))
    if not pvc:
        sys.exit('nenhum Kaizen da guia %s na faixa %d-%d' % (GUIA_PVC, ini, fim))
    ids = [int(r[0]) for r in pvc]
    if len(set(ids)) != len(ids):
        sys.exit('ID_KAIZEN duplicado na guia principal')

    idset = set(ids)
    aux_rows = {}
    for guia, _dest, cols, _decl, _lote in AUX:
        rs = ler(wb, guia, len(cols), 1)
        aux_rows[guia] = [r for r in rs if int(r[0]) in idset]

    # larguras exigidas, medidas no dado desta faixa
    req = []
    for i, (nome, _stg, kind) in enumerate(PVC_COLS):
        if kind == 'txt' and not nome.startswith('ID_'):
            m = maxlen(pvc, i)
            if m:
                req.append((GUIA_PVC, nome, m))
    hie = aux_rows['KZN_HIST_KAIZEN_HIERARQUIA']
    for i in range(2, 10):
        m = maxlen(hie, i)
        if m:
            req.append(('KZN_HIST_KAIZEN_HIERARQUIA', 'NM_HIERARQUIA_N%d' % (i - 1), m))

    o = []
    w = o.append
    faixa = '%d a %d' % (ini, fim)
    w("""/* =====================================================================
   Carga do historico nas tabelas KZN_HIST_* — %d Kaizens (%s)
   Fonte: %s, apenas as guias
   KZN_HIST_*. Valores literais: sem dependencia de arquivo externo.

   Volume: %d Kaizens, %d membros, %d resultados, %d hierarquias,
   %d desperdicios. Auxiliares filtradas pela lista exata de Kaizens da
   guia principal, nao pela faixa (os IDs nao sao contiguos).

   ---------------------------------------------------------------------
   1) '#N/A' NAS COLUNAS DE ID (erro de VLOOKUP na planilha)

   ID_USUARIO_CADASTRO, ID_USUARIO_LIDER, ID_APROVADOR e
   ID_USUARIO_ATUALIZACAO podem trazer '#N/A'. Entram no staging como
   texto e sao convertidos por TRY_CONVERT: '#N/A' vira NULL.

   Em ID_APROVADOR (coluna NULL-avel) isso e aceitavel. Nas outras tres,
   que sao NOT NULL, o NULL derrubaria a carga: a E1b detecta e ABORTA
   listando as linhas. Para carregar o resto sem corrigir a origem,
   troque @PULAR_INVALIDAS para 1 — as linhas problematicas (e suas
   auxiliares) ficam de fora e sao listadas.

   Nesta faixa: %s

   2) TRUNCAMENTO

   A E0.1 compara o tamanho DECLARADO de cada coluna com o exigido pelo
   dado desta faixa e ABORTA listando as que nao comportam — o INSERT so
   avisa com Msg 8152 se ANSI_WARNINGS estiver ligado; desligado, corta
   em silencio.

   2b) COLUNA NOT NULL QUE ESTA CARGA NAO PREENCHE

   criar_tabelas_hist.sql copia tipo e nulidade da tabela de producao, mas
   nao copia constraint DEFAULT. A E0.2 lista qualquer coluna NOT NULL sem
   DEFAULT fora do INSERT e aborta antes de gravar nada — caso tipico:
   SG_GM e ID_TIPO_KAIZEN, resolvidos por alinhar_hist_com_pvc_campos_gm.sql.

   3) ID_REPLICACAO vem como texto e e resolvido por lookup em
   CI.KZN_REPLICACAO; o que nao casar entra NULL e e listado na E7.

   Colunas de texto sao emitidas SEMPRE como literal string: num VALUES
   multi-linha o int tem precedencia sobre o nvarchar, e uma coluna com
   valores mistos faria o SQL Server tentar converter o texto para int.

   ---------------------------------------------------------------------
   A coluna de data da principal e resolvida em tempo de execucao:
   DT_CRIACAO se o renomear_dt_atualizacao_hist.sql ja rodou, senao
   DT_ATUALIZACAO.

   Idempotente (aborta se ja houver Kaizen na faixa). Transacionado.
   So escreve em CI.KZN_HIST_*. Schema: 'ci'.
   ===================================================================== */

SET NOCOUNT ON;
SET XACT_ABORT ON;

DECLARE @PULAR_INVALIDAS BIT = 0;   -- 1 = exclui linhas sem ID obrigatorio em vez de abortar
DECLARE @dtCol SYSNAME, @sql NVARCHAR(MAX), @qt INT;

/* =====================================================================
   E0 - PRE-CHECAGENS
   ===================================================================== */
DECLARE @tabs TABLE (TABELA SYSNAME PRIMARY KEY);
INSERT INTO @tabs (TABELA) VALUES
    (N'KZN_HIST_PEDRAVISAOCONSOLIDADA'), (N'KZN_HIST_MEMBROS_EQUIPE'),
    (N'KZN_HIST_RESULTADO_KAIZEN'),      (N'KZN_HIST_KAIZEN_HIERARQUIA'),
    (N'KZN_HIST_KAIZEN_DESPERDICIO');

IF EXISTS (SELECT 1 FROM @tabs WHERE OBJECT_ID('CI.' + TABELA, 'U') IS NULL)
BEGIN
    SELECT TABELA_AUSENTE = TABELA FROM @tabs WHERE OBJECT_ID('CI.' + TABELA, 'U') IS NULL;
    RAISERROR('Abortado: falta ao menos uma tabela CI.KZN_HIST_* (lista acima). Rode criar_tabelas_hist.sql antes.', 16, 1);
    RETURN;
END

IF EXISTS (SELECT 1 FROM CI.KZN_HIST_PEDRAVISAOCONSOLIDADA WHERE ID_KAIZEN BETWEEN %d AND %d)
BEGIN
    RAISERROR('Abortado: ja existem Kaizens na faixa %d-%d. Use apagar_dados_hist.sql antes de recarregar.', 16, 1);
    RETURN;
END

/* ---------------------------------------------------------------------
   E0.1 - TRUNCAMENTO: tamanho declarado x tamanho exigido por esta carga
   --------------------------------------------------------------------- */
DECLARE @req TABLE (TABELA SYSNAME, COLUNA SYSNAME, TAM_NECESSARIO INT);
INSERT INTO @req (TABELA, COLUNA, TAM_NECESSARIO) VALUES
%s;

/* max_length vem em BYTES: NVARCHAR gasta 2 por caractere, VARCHAR 1. */
IF EXISTS (SELECT 1 FROM @req r
           JOIN sys.columns c ON c.object_id = OBJECT_ID('CI.' + r.TABELA)
                             AND c.name COLLATE DATABASE_DEFAULT = r.COLUNA COLLATE DATABASE_DEFAULT
           JOIN sys.types ty ON ty.user_type_id = c.user_type_id
           WHERE c.max_length <> -1
             AND c.max_length / CASE WHEN ty.name IN ('nvarchar','nchar') THEN 2 ELSE 1 END < r.TAM_NECESSARIO)
BEGIN
    SELECT  TABELA = r.TABELA, COLUNA = c.name,
            TAM_DECLARADO  = c.max_length / CASE WHEN ty.name IN ('nvarchar','nchar') THEN 2 ELSE 1 END,
            TAM_NECESSARIO = r.TAM_NECESSARIO
    FROM    @req r
    JOIN    sys.columns c ON c.object_id = OBJECT_ID('CI.' + r.TABELA)
                         AND c.name COLLATE DATABASE_DEFAULT = r.COLUNA COLLATE DATABASE_DEFAULT
    JOIN    sys.types ty ON ty.user_type_id = c.user_type_id
    WHERE   c.max_length <> -1
      AND   c.max_length / CASE WHEN ty.name IN ('nvarchar','nchar') THEN 2 ELSE 1 END < r.TAM_NECESSARIO
    ORDER BY r.TABELA, c.name;
    RAISERROR('Abortado: ha coluna menor que o dado desta carga (lista acima). Amplie-a ou recrie as HIST com @TAM_TEXTO = -1 (MAX).', 16, 1);
    RETURN;
END

SET @dtCol = CASE WHEN COL_LENGTH('CI.KZN_HIST_PEDRAVISAOCONSOLIDADA','DT_CRIACAO') IS NOT NULL
                  THEN 'DT_CRIACAO' ELSE 'DT_ATUALIZACAO' END;
PRINT 'E0 - coluna de data da principal: ' + @dtCol;
""" % (len(pvc), faixa, src.split('/')[-1],
       len(pvc), len(aux_rows['KZN_HIST_MEMBROS_EQUIPE']),
       len(aux_rows['KZN_HIST_RESULTADO_KAIZEN']), len(hie),
       len(aux_rows['KZN_HIST_KAIZEN_DESPERDICIO']),
       _resumo_na(pvc),
       ini, fim, ini, fim,
       ',\n'.join("    (N'%s', N'%s', %d)" % r for r in req)))

    # guarda de ID_CATEGORIA: so faz sentido se a faixa tiver linha vazia
    n_cat = sum(1 for r in pvc if r[4] is None)
    if n_cat:
        w("""
IF EXISTS (SELECT 1 FROM sys.columns
           WHERE object_id = OBJECT_ID('CI.KZN_HIST_PEDRAVISAOCONSOLIDADA')
             AND name = 'ID_CATEGORIA' AND is_nullable = 0)
BEGIN
    RAISERROR('Abortado: ID_CATEGORIA e NOT NULL mas vem vazia em %d das %d linhas. Rode relaxar_not_null_hist.sql antes.', 16, 1);
    RETURN;
END
""" % (n_cat, len(pvc)))

    w("""
/* ---------------------------------------------------------------------
   E0.2 - NOT NULL sem DEFAULT que esta carga NAO preenche
   --------------------------------------------------------------------- */
DECLARE @carregadas TABLE (TABELA SYSNAME, COLUNA SYSNAME);
INSERT INTO @carregadas (TABELA, COLUNA)
SELECT N'KZN_HIST_PEDRAVISAOCONSOLIDADA', v FROM (VALUES
%s) x(v)
UNION ALL SELECT N'KZN_HIST_PEDRAVISAOCONSOLIDADA', @dtCol
%s;

IF EXISTS (
    SELECT 1
    FROM   @tabs t
    JOIN   sys.columns c ON c.object_id = OBJECT_ID('CI.' + t.TABELA)
    WHERE  c.is_nullable = 0 AND c.is_computed = 0 AND c.default_object_id = 0
      AND  NOT EXISTS (SELECT 1 FROM @carregadas g
                       WHERE g.TABELA COLLATE DATABASE_DEFAULT = t.TABELA COLLATE DATABASE_DEFAULT
                         AND g.COLUNA COLLATE DATABASE_DEFAULT = c.name COLLATE DATABASE_DEFAULT))
BEGIN
    SELECT  TABELA = t.TABELA, COLUNA = c.name, TIPO = ty.name
    FROM    @tabs t
    JOIN    sys.columns c  ON c.object_id = OBJECT_ID('CI.' + t.TABELA)
    JOIN    sys.types   ty ON ty.user_type_id = c.user_type_id
    WHERE   c.is_nullable = 0 AND c.is_computed = 0 AND c.default_object_id = 0
      AND   NOT EXISTS (SELECT 1 FROM @carregadas g
                        WHERE g.TABELA COLLATE DATABASE_DEFAULT = t.TABELA COLLATE DATABASE_DEFAULT
                          AND g.COLUNA COLLATE DATABASE_DEFAULT = c.name COLLATE DATABASE_DEFAULT)
    ORDER BY t.TABELA, c.column_id;
    RAISERROR('Abortado: ha coluna NOT NULL sem DEFAULT que esta carga nao preenche (lista acima) — o INSERT falharia com Msg 515. Para SG_GM/ID_TIPO_KAIZEN, rode alinhar_hist_com_pvc_campos_gm.sql, que cria os DEFAULT.', 16, 1);
    RETURN;
END
PRINT 'E0 ok - larguras e colunas obrigatorias conferidas.';

/* Sem GO ate o fim da carga: RAISERROR + RETURN so encerram o batch em
   que aparecem, e um GO aqui deixaria a carga rodar apos um aborto. */

BEGIN TRANSACTION;
BEGIN TRY

/* =====================================================================
   E1 - STAGING (IDs com '#N/A' e ID_REPLICACAO entram como texto)
   ===================================================================== */
CREATE TABLE #PVC (
%s);
""" % (',\n'.join('    (N\'%s\')' % c for c in PVC_DESTINO),
       '\n'.join("UNION ALL SELECT N'%s', v FROM (VALUES\n%s) x(v)"
                 % (dest, ',\n'.join('    (N\'%s\')' % c for c in cols))
                 for _g, dest, cols, _d, _l in AUX),
       STAGING_PVC))

    for r in pvc:
        vals = [lit(r[i], PVC_COLS[i][2]) for i in range(24)]
        w('INSERT INTO #PVC VALUES (%s);\n' % ', '.join(vals))

    w("""
/* =====================================================================
   E1b - TRAVA: ID obrigatorio que nao converte
   ===================================================================== */
IF EXISTS (SELECT 1 FROM #PVC
           WHERE TRY_CONVERT(INT, TX_CADASTRO) IS NULL
              OR TRY_CONVERT(INT, TX_LIDER) IS NULL
              OR TRY_CONVERT(INT, TX_ATUALIZACAO) IS NULL)
BEGIN
    SELECT  PROBLEMA = 'ID obrigatorio nao numerico',
            ID_KAIZEN, CADASTRO = TX_CADASTRO, LIDER = TX_LIDER, ATUALIZACAO = TX_ATUALIZACAO
    FROM    #PVC
    WHERE   TRY_CONVERT(INT, TX_CADASTRO) IS NULL
       OR   TRY_CONVERT(INT, TX_LIDER) IS NULL
       OR   TRY_CONVERT(INT, TX_ATUALIZACAO) IS NULL
    ORDER BY ID_KAIZEN;

    IF @PULAR_INVALIDAS = 0
        RAISERROR('Abortado: ha Kaizen(s) com ID obrigatorio nao numerico (lista acima). Corrija a origem, ou troque @PULAR_INVALIDAS para 1 para carregar o restante.', 16, 1);
    ELSE
    BEGIN
        DELETE FROM #PVC
        WHERE TRY_CONVERT(INT, TX_CADASTRO) IS NULL
           OR TRY_CONVERT(INT, TX_LIDER) IS NULL
           OR TRY_CONVERT(INT, TX_ATUALIZACAO) IS NULL;
        PRINT '  E1b - ' + CAST(@@ROWCOUNT AS VARCHAR(10)) + ' Kaizen(s) excluido(s) da carga.';
    END
END

/* =====================================================================
   E2 - Tabela principal
   ===================================================================== */
SET @sql = N'
INSERT INTO CI.KZN_HIST_PEDRAVISAOCONSOLIDADA
    (ID_KAIZEN, ID_USUARIO_CADASTRO, ID_USUARIO_LIDER, NM_KAIZEN, ID_CATEGORIA, ID_REPLICACAO,
     DS_PROBLEMA, DS_OBJETIVO, ID_APROVADOR, URL_IMG_ANTES, DS_ESTADO_ANTES, URL_IMG_DEPOIS,
     DS_ESTADO_DEPOIS, URL_REFERENCIA, DS_COMPARA_META, DS_LICOES_APRENDIDAS,
     VL_RESULTADO_FINANCEIRO, ID_MOEDA, DS_RESULTADO_ALCANCADO, DT_CONCLUSAO, ' + QUOTENAME(@dtCol) + N',
     ID_USUARIO_ATUALIZACAO, DS_MOTIVO, ID_STATUS)
SELECT  p.ID_KAIZEN, TRY_CONVERT(INT, p.TX_CADASTRO), TRY_CONVERT(INT, p.TX_LIDER), p.NM_KAIZEN, p.ID_CATEGORIA, r.ID_REPLICACAO,
        p.DS_PROBLEMA, p.DS_OBJETIVO, TRY_CONVERT(INT, p.TX_APROVADOR), p.URL_IMG_ANTES, p.DS_ESTADO_ANTES, p.URL_IMG_DEPOIS,
        p.DS_ESTADO_DEPOIS, p.URL_REFERENCIA, p.DS_COMPARA_META, p.DS_LICOES_APRENDIDAS,
        p.VL_RESULTADO_FINANCEIRO, p.ID_MOEDA, p.DS_RESULTADO_ALCANCADO, p.DT_CONCLUSAO, p.DT_REG,
        TRY_CONVERT(INT, p.TX_ATUALIZACAO), p.DS_MOTIVO, p.ID_STATUS
FROM        #PVC p
OUTER APPLY (SELECT TOP (1) d.ID_REPLICACAO
             FROM   CI.KZN_REPLICACAO d
             WHERE  LTRIM(RTRIM(d.NM_REPLICACAO)) COLLATE Latin1_General_CI_AI
                    IN (LTRIM(RTRIM(p.TX_REPLICACAO))                                                COLLATE Latin1_General_CI_AI,
                        LTRIM(RTRIM(LEFT(p.TX_REPLICACAO, NULLIF(CHARINDEX(''/'', p.TX_REPLICACAO),0) - 1)))      COLLATE Latin1_General_CI_AI,
                        LTRIM(RTRIM(SUBSTRING(p.TX_REPLICACAO, NULLIF(CHARINDEX(''/'', p.TX_REPLICACAO),0) + 1, 200))) COLLATE Latin1_General_CI_AI)
             ORDER BY d.ID_IDIOMA) r
ORDER BY p.ID_KAIZEN;
SET @c = @@ROWCOUNT;';
/* @c por OUTPUT, medido DENTRO do batch dinamico: ler @@ROWCOUNT aqui
   fora, depois do EXEC, e fragil. */
EXEC sp_executesql @sql, N'@c INT OUTPUT', @c = @qt OUTPUT;
PRINT '  E2 - KZN_HIST_PEDRAVISAOCONSOLIDADA: ' + CAST(@qt AS VARCHAR(10)) + ' linha(s).';
""")

    rotulos = {'KZN_HIST_MEMBROS_EQUIPE': 'E3 - Membros de equipe',
               'KZN_HIST_RESULTADO_KAIZEN': 'E4 - Resultados',
               'KZN_HIST_KAIZEN_HIERARQUIA': 'E5 - Fotografia da hierarquia',
               'KZN_HIST_KAIZEN_DESPERDICIO': 'E6 - Desperdicios'}
    for guia, dest, cols, decl, lote in AUX:
        rs = aux_rows[guia]
        kinds = ['int', 'int'] + ['txt'] * (len(cols) - 3) + ['dt'] if len(cols) > 3 else ['int', 'int', 'dt']
        w("""
/* =====================================================================
   %s
   ===================================================================== */
CREATE TABLE #T (%s);
""" % (rotulos[guia], decl))
        for k in range(0, len(rs), lote):
            bloco = rs[k:k + lote]
            w('INSERT INTO #T VALUES\n')
            w(',\n'.join('(%s)' % ', '.join(lit(r[i], kinds[i]) for i in range(len(cols)))
                         for r in bloco) + ';\n')
        w("""INSERT INTO CI.%s (%s)
SELECT %s FROM #T t
WHERE EXISTS (SELECT 1 FROM #PVC p WHERE p.ID_KAIZEN = t.ID_KAIZEN);
PRINT '  %s: ' + CAST(@@ROWCOUNT AS VARCHAR(10)) + ' linha(s).';
DROP TABLE #T;
""" % (dest, ', '.join(cols), ', '.join('t.' + c for c in cols), dest))

    w("""
/* =====================================================================
   E7 - Pendencias (antes do COMMIT)
   ===================================================================== */
SELECT  PROBLEMA = 'REPLICACAO sem correspondencia em CI.KZN_REPLICACAO',
        TEXTO_DA_PLANILHA = p.TX_REPLICACAO, QT_KAIZENS = COUNT(*)
FROM    #PVC p
JOIN    CI.KZN_HIST_PEDRAVISAOCONSOLIDADA h ON h.ID_KAIZEN = p.ID_KAIZEN
WHERE   h.ID_REPLICACAO IS NULL AND p.TX_REPLICACAO IS NOT NULL
GROUP BY p.TX_REPLICACAO;

SELECT  PROBLEMA = 'ID_APROVADOR nao numerico -> gravado NULL',
        QT_KAIZENS = COUNT(*)
FROM    #PVC WHERE TX_APROVADOR IS NOT NULL AND TRY_CONVERT(INT, TX_APROVADOR) IS NULL;

DROP TABLE #PVC;

COMMIT TRANSACTION;
PRINT 'Carga concluida.';

END TRY
BEGIN CATCH
    IF XACT_STATE() <> 0 ROLLBACK TRANSACTION;
    PRINT 'ERRO - nada foi gravado (rollback aplicado): ' + ERROR_MESSAGE();
    THROW;
END CATCH
GO

/* =====================================================================
   E8 - CONFERENCIA
   ===================================================================== */
SELECT TABELA='KZN_HIST_PEDRAVISAOCONSOLIDADA', LINHAS=COUNT(*) FROM CI.KZN_HIST_PEDRAVISAOCONSOLIDADA WHERE ID_KAIZEN BETWEEN %d AND %d
UNION ALL SELECT 'KZN_HIST_MEMBROS_EQUIPE',     COUNT(*) FROM CI.KZN_HIST_MEMBROS_EQUIPE      WHERE ID_KAIZEN BETWEEN %d AND %d
UNION ALL SELECT 'KZN_HIST_RESULTADO_KAIZEN',   COUNT(*) FROM CI.KZN_HIST_RESULTADO_KAIZEN    WHERE ID_KAIZEN BETWEEN %d AND %d
UNION ALL SELECT 'KZN_HIST_KAIZEN_HIERARQUIA',  COUNT(*) FROM CI.KZN_HIST_KAIZEN_HIERARQUIA   WHERE ID_KAIZEN BETWEEN %d AND %d
UNION ALL SELECT 'KZN_HIST_KAIZEN_DESPERDICIO', COUNT(*) FROM CI.KZN_HIST_KAIZEN_DESPERDICIO  WHERE ID_KAIZEN BETWEEN %d AND %d;
GO
""" % ((ini, fim) * 5))

    open(out, 'w', encoding='utf-8').write(''.join(o))
    print(len(pvc), len(aux_rows['KZN_HIST_MEMBROS_EQUIPE']),
          len(aux_rows['KZN_HIST_RESULTADO_KAIZEN']), len(hie),
          len(aux_rows['KZN_HIST_KAIZEN_DESPERDICIO']))


def _resumo_na(pvc):
    idx = {'ID_USUARIO_CADASTRO': 1, 'ID_USUARIO_LIDER': 2, 'ID_APROVADOR': 8,
           'ID_USUARIO_ATUALIZACAO': 21}
    partes = []
    for nome, i in idx.items():
        n = sum(1 for r in pvc if isinstance(r[i], str) and r[i].strip() == '#N/A')
        if n:
            partes.append('%s em %d linha(s)' % (nome, n))
    return '; '.join(partes) if partes else 'nenhum \'#N/A\' — a E1b nao deve disparar.'


if __name__ == '__main__':
    main()
